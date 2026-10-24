import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from src.purchase_processor import build_summary
from src.unloading_slip import load_combined_purchase_rows


class SchemePurchaseSummaryTests(unittest.TestCase):
    def _source(self, path: Path, unit: str = "CAR") -> None:
        wb = Workbook()
        ws = wb.active
        ws.title = "Invoice"
        headers = {
            8: "ITEM",
            21: "MRP",
            24: "PTR",
            25: "Case",
            27: "QTY.",
            31: "Billing Rate",
            33: "AMOUNT",
            35: "DISCOUNT",
            37: "TAXABLE AMOUNT",
            40: "CGST Rate%",
            44: "SGST Rate%",
            47: "IGST Rate%",
            52: "TOTAL AMOUNT",
            54: "Invoice No.",
            57: "Date",
        }
        for column, value in headers.items():
            ws.cell(2, column, value)
        paid = {
            8: "DOUBLE ENGINE DETERGENT POWDER 1KG",
            21: 80,
            24: 65,
            25: 105,
            27: 3150,
            31: 50.46,
            33: 158949,
            35: 0,
            37: 158949,
            40: 9,
            44: 9,
            47: 0,
            52: 187559.82,
            54: "SG-DT-26-01750",
            57: "22/10/2026",
        }
        for column, value in paid.items():
            ws.cell(3, column, value)
        ws.cell(8, 3, "Scheme Item")
        scheme = {
            3: 2,
            6: "SCH DOUBLE ENGINE DETERGENT POWDER 1KG*",
            11: "34021200",
            18: 80,
            20: 1950,
            22: 7,
            23: 7,
            26: unit,
            28: 1513.69,
            30: 10595.83,
            32: 10595.83,
            34: 0,
            36: 9,
            39: 9,
            42: 0,
            46: 0,
        }
        for column, value in scheme.items():
            ws.cell(10, column, value)
        wb.save(path)

    @patch("src.tally_item_master.get_item_master")
    @patch("src.tally_scheme_items.resolve_scheme_items")
    def test_scheme_item_reaches_summary_and_unloading(self, resolve, item_master):
        resolve.return_value = {
            "SCH DOUBLE ENGINE DETERGENT POWDER 1KG*": {
                "pcs_per_case": 30,
                "combinations": [("65.00", 80, 51.97)],
            }
        }
        item_master.return_value = {
            "items": [
                "DOUBLE ENGINE DETERGENT POWDER 1KG",
                "SCH DOUBLE ENGINE DETERGENT POWDER 1KG*",
            ],
            "source": "test",
            "fetched_at": "2026-10-24T00:00:00",
            "company_fingerprint": "test",
            "warning": "",
        }
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / "invoice.xlsx"
            output = Path(folder) / "summary.xlsx"
            self._source(source)
            result = build_summary(source, output)
            self.assertEqual(result["output_rows"], 2)
            self.assertEqual(result["total_case"], 112)
            self.assertEqual(result["total_qty"], 3360)
            self.assertEqual(result["scheme_confirmations"], ["SCHEME ITEM - CONFIRM PTR 65.00"])

            wb = load_workbook(output, data_only=True)
            ws = wb["Purchase Summary"]
            values = [ws.cell(3, column).value for column in range(1, 13)]
            self.assertEqual(values[1:9], ["65.00", 80, 51.97, "SCH DOUBLE ENGINE DETERGENT POWDER 1KG*", 7, 210, 50.46, 1])
            self.assertEqual(values[11], "SCHEME ITEM - CONFIRM PTR 65.00")

            rows = load_combined_purchase_rows([output], datetime(2026, 10, 22))
            scheme = next(row for row in rows if row["item"].startswith("SCH DOUBLE"))
            self.assertEqual(scheme["cases"], 7)
            self.assertEqual(scheme["pcs"], 210)

    @patch("src.tally_item_master.get_item_master")
    @patch("src.tally_scheme_items.resolve_scheme_items")
    def test_unknown_scheme_unit_is_blocked(self, resolve, item_master):
        resolve.return_value = {
            "SCH DOUBLE ENGINE DETERGENT POWDER 1KG*": {
                "pcs_per_case": 30,
                "combinations": [("65.00", 80, 51.97)],
            }
        }
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / "invoice.xlsx"
            self._source(source, unit="BOX")
            with self.assertRaisesRegex(ValueError, "unsupported unit"):
                build_summary(source, Path(folder) / "summary.xlsx")


if __name__ == "__main__":
    unittest.main()
