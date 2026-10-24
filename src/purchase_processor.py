from __future__ import annotations

import argparse
import re
from collections import OrderedDict
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


REQUIRED_HEADERS = {
    "item": "ITEM",
    "ptr": "PTR",
    "mrp": "MRP",
    "case": "Case",
    "qty": "QTY.",
    "billing_rate": "Billing Rate",
    "amount": "AMOUNT",
    "discount": "DISCOUNT",
    "invoice_no": "Invoice No.",
    "date": "Date",
}


def _clean_header(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip()).casefold()


def _clean_filename_part(value: Any) -> str:
    text = str(value or "").strip()
    text = text.replace("/", "-").replace("\\", "-")
    text = re.sub(r'[<>:"|?*]+', "-", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip(" .-")


def _as_float(value: Any) -> float:
    if value is None or value == "":
        return 0.0
    return float(value)


def _find_header_row_and_columns(ws) -> tuple[int, dict[str, int]]:
    wanted = {_clean_header(label): key for key, label in REQUIRED_HEADERS.items()}
    for row in range(1, min(ws.max_row or 1, 20) + 1):
        found: dict[str, int] = {}
        for col in range(1, (ws.max_column or 1) + 1):
            header = _clean_header(ws.cell(row, col).value)
            if header in wanted:
                found[wanted[header]] = col
        if set(REQUIRED_HEADERS) <= set(found):
            return row, found
    missing = ", ".join(REQUIRED_HEADERS.values())
    raise ValueError(f"Could not find required invoice headers: {missing}")


def _is_item_row(item: str) -> bool:
    normalized = item.strip().casefold()
    return bool(normalized) and normalized != "total"


def _scheme_source_rows(ws) -> list[dict[str, Any]]:
    markers: list[tuple[int, int]] = []
    for row in range(1, (ws.max_row or 1) + 1):
        for column in range(1, (ws.max_column or 1) + 1):
            if _clean_header(ws.cell(row, column).value) == "scheme item":
                markers.append((row, column))

    rows: list[dict[str, Any]] = []
    for marker_index, (marker_row, marker_col) in enumerate(markers):
        end_row = markers[marker_index + 1][0] - 1 if marker_index + 1 < len(markers) else ws.max_row
        for row in range(marker_row + 1, end_row + 1):
            item = str(ws.cell(row, marker_col + 3).value or "").strip()
            if not _is_item_row(item):
                continue
            rows.append(
                {
                    "item": item,
                    "hsn": str(ws.cell(row, marker_col + 8).value or "").strip(),
                    "invoice_mrp": ws.cell(row, marker_col + 15).value,
                    "case_ptr": _as_float(ws.cell(row, marker_col + 17).value),
                    "case": _as_float(ws.cell(row, marker_col + 19).value),
                    "case_qty": _as_float(ws.cell(row, marker_col + 20).value),
                    "unit": str(ws.cell(row, marker_col + 23).value or "").strip(),
                    "case_billing_rate": _as_float(ws.cell(row, marker_col + 25).value),
                    "amount": _as_float(ws.cell(row, marker_col + 27).value),
                    "discount": _as_float(ws.cell(row, marker_col + 29).value),
                }
            )
    return rows


def _scheme_records(source_rows: list[dict[str, Any]], invoice_no: str, invoice_date: str) -> list[dict[str, Any]]:
    if not source_rows:
        return []
    from .tally_scheme_items import resolve_scheme_items

    details = resolve_scheme_items([row["item"] for row in source_rows])
    grouped: OrderedDict[tuple[str, str, str, str, str], dict[str, Any]] = OrderedDict()
    for source in source_rows:
        item = source["item"]
        if source["unit"].casefold() != "car":
            raise ValueError(
                f"Scheme item {item!r} uses unsupported unit {source['unit']!r}; expected CAR (Case)."
            )
        if source["case"] <= 0 or source["case_qty"] <= 0:
            raise ValueError(f"Scheme item {item!r} has no valid Case quantity in the invoice.")
        if abs(source["case"] - source["case_qty"]) > 0.00001:
            raise ValueError(
                f"Scheme item {item!r} has inconsistent Case and CAR quantities: "
                f"{source['case']} and {source['case_qty']}."
            )
        item_details = details[item]
        conversion = Decimal(str(item_details["pcs_per_case"]))
        case_ptr = Decimal(str(source["case_ptr"]))
        case_rate = Decimal(str(source["case_billing_rate"]))
        if case_ptr <= 0 or case_rate < 0:
            raise ValueError(f"Scheme item {item!r} has invalid PTR or Billing Rate in the invoice.")
        ptr_per_piece = (case_ptr / conversion).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        billing_rate = (case_rate / conversion).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        matching = [
            combination
            for combination in item_details["combinations"]
            if Decimal(combination[0]) == ptr_per_piece
        ]
        matching = list(dict.fromkeys(matching))
        if not matching:
            raise ValueError(
                f"Scheme item {item!r}: calculated PTR {ptr_per_piece:.2f} per PCS does not exist "
                "in Tally. Ask the user for PTR, MRP, Sell Rate and Purchase Rate per PCS."
            )
        if len(matching) > 1:
            options = ", ".join(
                f"PTR {ptr} / MRP {mrp:.2f} / Sell {sell:.2f}" for ptr, mrp, sell in matching
            )
            raise ValueError(
                f"Scheme item {item!r}: multiple Tally combinations match calculated PTR "
                f"{ptr_per_piece:.2f}: {options}. Ask the user to select one."
            )
        ptr, mrp, sell_rate = matching[0]
        key = (item, ptr, str(mrp), str(sell_rate), str(billing_rate))
        if key not in grouped:
            grouped[key] = {
                "ptr": ptr,
                "mrp": float(mrp),
                "sell_rate": float(sell_rate),
                "item": item,
                "case": 0.0,
                "qty": 0.0,
                "billing_rate": float(billing_rate),
                "discount": 0.0,
                "amount": 0.0,
                "invoice_no": invoice_no,
                "date": invoice_date,
                "status": f"SCHEME ITEM - CONFIRM PTR {ptr}",
                "is_scheme": True,
            }
        record = grouped[key]
        record["case"] += source["case"]
        record["qty"] += float(Decimal(str(source["case_qty"])) * conversion)
        record["amount"] += source["amount"]
        record["discount"] += source["discount"]
    return list(grouped.values())


def build_summary(source_path: Path, output_path: Path | None = None) -> dict[str, Any]:
    source_path = Path(source_path)
    if not source_path.exists():
        raise FileNotFoundError(source_path)

    raw_values_wb = load_workbook(source_path, data_only=True)
    raw_values_ws = raw_values_wb[raw_values_wb.sheetnames[0]]
    header_row, columns = _find_header_row_and_columns(raw_values_ws)

    groups: OrderedDict[tuple[str, float], dict[str, Any]] = OrderedDict()
    input_rows = 0
    first_invoice_no = ""
    first_invoice_date = ""

    for row in range(header_row + 1, (raw_values_ws.max_row or header_row) + 1):
        item = str(raw_values_ws.cell(row, columns["item"]).value or "").strip()
        if not _is_item_row(item):
            continue

        input_rows += 1
        billing_rate = _as_float(raw_values_ws.cell(row, columns["billing_rate"]).value)
        key = (item, billing_rate)

        invoice_no = str(raw_values_ws.cell(row, columns["invoice_no"]).value or "").strip()
        invoice_date = str(raw_values_ws.cell(row, columns["date"]).value or "").strip()
        first_invoice_no = first_invoice_no or invoice_no
        first_invoice_date = first_invoice_date or invoice_date

        if key not in groups:
            groups[key] = {
                "ptr": _as_float(raw_values_ws.cell(row, columns["ptr"]).value),
                "mrp": _as_float(raw_values_ws.cell(row, columns["mrp"]).value),
                "item": item,
                "case": 0.0,
                "qty": 0.0,
                "billing_rate": billing_rate,
                "discount": 0.0,
                "amount": 0.0,
                "invoice_no": invoice_no,
                "date": invoice_date,
            }

        group = groups[key]
        group["case"] += _as_float(raw_values_ws.cell(row, columns["case"]).value)
        group["qty"] += _as_float(raw_values_ws.cell(row, columns["qty"]).value)
        group["discount"] += _as_float(raw_values_ws.cell(row, columns["discount"]).value)
        group["amount"] += _as_float(raw_values_ws.cell(row, columns["amount"]).value)

    scheme_sources = _scheme_source_rows(raw_values_ws)
    input_rows += len(scheme_sources)
    records = list(groups.values())
    records.extend(_scheme_records(scheme_sources, first_invoice_no, first_invoice_date))
    records = sorted(records, key=lambda record: record["item"].casefold())

    if output_path is None:
        output_path = source_path.with_name(_filename_from_invoice(source_path, first_invoice_no, first_invoice_date))
    output_path = Path(output_path)

    from .tally_item_master import get_item_master

    item_master = get_item_master()
    tally_items = set(item_master["items"])
    new_items = sorted({record["item"] for record in records if record["item"] not in tally_items})

    output_wb = load_workbook(source_path)
    raw_ws = output_wb[output_wb.sheetnames[0]]
    raw_ws.title = "Raw Data"
    if "Purchase Summary" in output_wb.sheetnames:
        del output_wb["Purchase Summary"]
    output_ws = output_wb.create_sheet("Purchase Summary")

    headers = [
        "No.",
        "PTR",
        "MRP",
        "Sell Rate",
        "ITEM",
        "Case",
        "QTY",
        "Billing Rate",
        "Discount %",
        "Invoice No.",
        "Date",
        "Status",
    ]
    output_ws.append(headers)

    for serial, record in enumerate(records, start=1):
        multiplier = 1.025 if record["item"].casefold().startswith("tea") else 1.03
        sell_rate = record.get("sell_rate", round(record["billing_rate"] * multiplier, 2))
        discount_percent = (
            round((record["discount"] / record["amount"]) * 100, 2) / 100
            if record["amount"]
            else 0
        )
        output_ws.append(
            [
                serial,
                record["ptr"],
                record["mrp"],
                sell_rate,
                record["item"],
                record["case"],
                record["qty"],
                record["billing_rate"],
                discount_percent,
                record["invoice_no"],
                record["date"],
                "NEW ITEM" if record["item"] in new_items else record.get("status", ""),
            ]
        )

    _format_summary_sheet(output_ws, new_items)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_wb.save(output_path)

    return {
        "source": str(source_path),
        "output": str(output_path),
        "invoice_no": first_invoice_no,
        "invoice_date": first_invoice_date,
        "input_rows": input_rows,
        "output_rows": len(records),
        "consolidated_rows": input_rows - len(records),
        "total_case": sum(record["case"] for record in records),
        "total_qty": sum(record["qty"] for record in records),
        "created_at": datetime.now().isoformat(timespec="seconds"),
        "new_items": new_items,
        "item_master_source": item_master["source"],
        "item_master_fetched_at": item_master["fetched_at"],
        "item_master_company_fingerprint": item_master["company_fingerprint"],
        "item_master_warning": item_master.get("warning", ""),
        "scheme_confirmations": [
            record["status"] for record in records if record.get("is_scheme") and record.get("status")
        ],
    }


def default_output_filename(source_path: Path) -> str:
    source_path = Path(source_path)
    values_wb = load_workbook(source_path, data_only=True, read_only=False)
    values_ws = values_wb[values_wb.sheetnames[0]]
    header_row, columns = _find_header_row_and_columns(values_ws)
    for row in range(header_row + 1, (values_ws.max_row or header_row) + 1):
        item = str(values_ws.cell(row, columns["item"]).value or "").strip()
        if _is_item_row(item):
            invoice_no = str(values_ws.cell(row, columns["invoice_no"]).value or "").strip()
            invoice_date = str(values_ws.cell(row, columns["date"]).value or "").strip()
            return _filename_from_invoice(source_path, invoice_no, invoice_date)
    return f"{_clean_filename_part(source_path.stem)} Summary.xlsx"


def _filename_from_invoice(source_path: Path, invoice_no: str, invoice_date: str) -> str:
    safe_invoice = _clean_filename_part(invoice_no) or source_path.stem
    safe_date = _clean_filename_part(invoice_date)
    return f"{safe_invoice} {safe_date} Summary.xlsx".strip()


def _format_summary_sheet(ws, new_items: list[str]) -> None:
    max_row = ws.max_row
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in range(2, max_row + 1):
        ws.cell(row, 1).alignment = Alignment(horizontal="center")
        ws.cell(row, 10).alignment = Alignment(horizontal="center")
        ws.cell(row, 11).alignment = Alignment(horizontal="center")
        for column in (2, 3, 4, 8):
            ws.cell(row, column).number_format = "0.00"
        ws.cell(row, 6).number_format = "General"
        ws.cell(row, 7).number_format = "0"
        ws.cell(row, 9).number_format = "0.00%"
        if ws.cell(row, 5).value in new_items or ws.cell(row, 12).value:
            yellow = PatternFill("solid", fgColor="FFF2CC")
            for column in range(1, 13):
                ws.cell(row, column).fill = yellow
            ws.cell(row, 12).font = Font(bold=True, color="9C6500")

    widths = {
        "A": 7,
        "B": 10,
        "C": 10,
        "D": 12,
        "E": 48,
        "F": 12,
        "G": 12,
        "H": 14,
        "I": 14,
        "J": 20,
        "K": 14,
        "L": 14,
    }
    for column, width in widths.items():
        ws.column_dimensions[column].width = width

    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    ws.row_dimensions[1].height = 30

    if max_row > 1:
        table = Table(displayName="PurchaseSummaryTable", ref=f"A1:L{max_row}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        ws.add_table(table)


def main() -> None:
    parser = argparse.ArgumentParser(description="Build Goldiee purchase summary workbook.")
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path, nargs="?")
    args = parser.parse_args()
    result = build_summary(args.source, args.output)
    for key, value in result.items():
        print(f"{key}: {value}")


if __name__ == "__main__":
    main()
