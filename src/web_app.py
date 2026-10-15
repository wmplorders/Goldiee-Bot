"""Local Goldiee control panel. Jobs run serially in isolated subprocesses."""
import argparse
import json
import mimetypes
import os
import secrets
import subprocess
import sys
import threading
import uuid
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import urlparse, parse_qs, quote

from .settings import PROJECT_ROOT, load_config

ASSETS = Path(__file__).with_name('web')
JOBS_DIR = PROJECT_ROOT / 'data' / 'web_jobs'
CONVERSATIONS_DIR = PROJECT_ROOT / 'data' / 'web_conversations'
SETTINGS = PROJECT_ROOT / 'data' / 'web_settings.json'
POOL = ThreadPoolExecutor(max_workers=1)
LOCK = threading.Lock()
TOKEN = secrets.token_urlsafe(32)
PASSWORD = os.environ.get('GOLDIEE_WEB_PASSWORD', '')
SESSIONS = set()


def settings():
    return json.loads(SETTINGS.read_text()) if SETTINGS.exists() else {'company': load_config()['tally']['company_name']}


def save_job(job):
    JOBS_DIR.mkdir(parents=True, exist_ok=True)
    with LOCK:
        p = JOBS_DIR / (job['id'] + '.json')
        temp = p.with_suffix('.tmp')
        temp.write_text(json.dumps(job), encoding='utf-8')
        temp.replace(p)


def jobs():
    with LOCK:
        return sorted([json.loads(p.read_text()) for p in JOBS_DIR.glob('*.json')], key=lambda x:x['created'], reverse=True)[:100]


def conversations():
    CONVERSATIONS_DIR.mkdir(parents=True, exist_ok=True)
    with LOCK:
        result = []
        for path in CONVERSATIONS_DIR.glob('*.json'):
            try:
                result.append(json.loads(path.read_text(encoding='utf-8')))
            except (OSError, json.JSONDecodeError):
                continue
        return sorted(result, key=lambda x:x.get('updated', x.get('created','')), reverse=True)[:200]


def save_conversation(conversation):
    identifier = str(conversation.get('id',''))
    if not identifier or not all(c in '0123456789abcdef' for c in identifier) or len(identifier) != 32:
        raise ValueError('Invalid conversation')
    messages = conversation.get('messages', [])
    if not isinstance(messages, list) or len(messages) > 250:
        raise ValueError('Conversation is too large')
    encoded = json.dumps(conversation, ensure_ascii=False)
    if len(encoded) > 1_500_000:
        raise ValueError('Conversation is too large')
    CONVERSATIONS_DIR.mkdir(parents=True, exist_ok=True)
    with LOCK:
        path = CONVERSATIONS_DIR / (identifier + '.json')
        temp = path.with_suffix('.tmp')
        temp.write_text(encoded, encoding='utf-8')
        temp.replace(path)


def run_job(job):
    job.update(status='running', detail='Processing; detailed Tally exports can take several minutes.')
    save_job(job)
    env = os.environ.copy()
    env['TALLY_COMPANY_NAME'] = job['company']
    env['TALLY_CLOSE_OTHER_COMPANIES'] = 'false'
    try:
        r = subprocess.run([sys.executable, '-m', 'src.web_jobs', job['id']], cwd=PROJECT_ROOT,
                           env=env, capture_output=True, text=True, timeout=1800)
        if r.returncode:
            diagnostic = (r.stderr or r.stdout)[-6000:]
            job['diagnostic'] = diagnostic
            last_line = diagnostic.strip().splitlines()[-1]
            if ': ' in last_line:
                last_line = last_line.split(': ', 1)[1]
            raise RuntimeError(last_line)
        result = json.loads(r.stdout.strip().splitlines()[-1])
        job.update(status='completed', detail=result['message'], files=result.get('files', []))
    except Exception as exc:
        job.update(status='failed', detail=str(exc))
    job['finished'] = datetime.now().isoformat(timespec='seconds')
    save_job(job)


def file_path(name, folder='processed'):
    if Path(name).name != name or '/' in name or '\\' in name:
        raise ValueError('Invalid filename')
    root = Path(load_config()['paths'][folder]).resolve()
    p = (root / name).resolve()
    if p.parent != root or p.suffix.lower() not in {'.xlsx','.xlsm','.pdf'}:
        raise ValueError('Unsupported file')
    return p


class Handler(BaseHTTPRequestHandler):
    def reply(self, code, value, content_type='application/json'):
        data = json.dumps(value).encode() if content_type == 'application/json' else value
        self.send_response(code)
        self.send_header('Content-Type', content_type)
        self.send_header('Content-Length', str(len(data)))
        self.send_header('Cache-Control', 'no-store')
        self.send_header('X-Content-Type-Options', 'nosniff')
        self.end_headers()
        self.wfile.write(data)

    def authenticated(self):
        if not PASSWORD:
            return True
        cookie = self.headers.get('Cookie','')
        return any(('session=' + s) in cookie.split('; ') for s in SESSIONS)

    def do_GET(self):
        try:
            if not self.valid_host():
                return self.reply(403, {'error':'Host rejected'})
            path = urlparse(self.path).path
            if path in {'/', '/app.js', '/style.css'}:
                p = ASSETS / ('index.html' if path == '/' else path[1:])
                return self.reply(200, p.read_bytes(), mimetypes.guess_type(p)[0] or 'text/plain')
            if not self.authenticated():
                return self.reply(401, {'error':'Please sign in.'})
            if path == '/api/state':
                c = load_config()
                files = sorted(Path(c['paths']['processed']).glob('*'), key=lambda p:p.stat().st_mtime, reverse=True)
                return self.reply(200, dict(token=TOKEN, settings=settings(), jobs=jobs(), conversations=conversations(), files=[
                    dict(name=p.name, size=p.stat().st_size, modified=datetime.fromtimestamp(p.stat().st_mtime).isoformat(timespec='seconds'))
                    for p in files if p.is_file() and not p.name.startswith('~$') and p.suffix.lower() in {'.xlsx','.pdf'}][:300]))
            if path == '/api/companies':
                from .tally_context import list_loaded_companies
                return self.reply(200, {'companies':list_loaded_companies()})
            if path == '/api/file':
                name = parse_qs(urlparse(self.path).query)['name'][0]
                p = file_path(name)
                self.send_response(200)
                self.send_header('Content-Type', mimetypes.guess_type(p)[0] or 'application/octet-stream')
                self.send_header('Content-Disposition', "attachment; filename*=UTF-8''" + quote(p.name))
                self.end_headers()
                self.wfile.write(p.read_bytes())
                return
            self.reply(404, {'error':'Not found'})
        except Exception as exc:
            self.reply(400, {'error':str(exc)})

    def do_POST(self):
        try:
            if not self.valid_host():
                return self.reply(403, {'error':'Host rejected'})
            # Reject cross-origin submissions, including local-network CSRF.
            origin = self.headers.get('Origin')
            if origin and urlparse(origin).netloc != self.headers.get('Host'):
                return self.reply(403, {'error':'Origin rejected'})
            length = int(self.headers.get('Content-Length',0))
            if length > 30*1024*1024:
                return self.reply(413, {'error':'Maximum upload is 30 MB.'})
            body = self.rfile.read(length)
            if self.path == '/api/login':
                data = json.loads(body)
                if not PASSWORD or secrets.compare_digest(data.get('password',''), PASSWORD):
                    session = secrets.token_urlsafe(32)
                    SESSIONS.add(session)
                    self.send_response(200)
                    self.send_header('Set-Cookie', 'session='+session+'; HttpOnly; SameSite=Strict; Path=/')
                    self.end_headers()
                    self.wfile.write(b'{}')
                    return
                return self.reply(401, {'error':'Incorrect password'})
            if not self.authenticated() or self.headers.get('X-CSRF-Token') != TOKEN:
                return self.reply(403, {'error':'Session expired. Refresh the page.'})
            if self.path == '/api/upload':
                name = self.headers.get('X-Filename', 'invoice.xlsx')
                from urllib.parse import unquote
                name = uuid.uuid4().hex[:8] + '_' + unquote(name)
                p = file_path(name, 'uploads')
                if p.suffix.lower() not in {'.xlsx','.xlsm'}:
                    raise ValueError('Upload an Excel workbook')
                p.parent.mkdir(parents=True, exist_ok=True)
                p.write_bytes(body)
                return self.reply(200, {'name':name})
            data = json.loads(body)
            if self.path == '/api/conversations':
                if data.get('action') == 'create':
                    now = datetime.now().isoformat(timespec='seconds')
                    conversation = dict(
                        id=uuid.uuid4().hex,
                        title=str(data.get('title') or 'New conversation')[:120],
                        flow=str(data.get('flow') or '')[:40],
                        status='active',
                        created=now,
                        updated=now,
                        messages=data.get('messages', [])[:250],
                        context=data.get('context', {}),
                    )
                    save_conversation(conversation)
                    return self.reply(201, conversation)
                if data.get('action') == 'save':
                    conversation = data.get('conversation', {})
                    conversation['updated'] = datetime.now().isoformat(timespec='seconds')
                    save_conversation(conversation)
                    return self.reply(200, conversation)
                raise ValueError('Unknown conversation action')
            if self.path == '/api/preview/unloading':
                from .unloading_slip import parse_user_date, find_purchase_summaries, same_date
                from openpyxl import load_workbook
                target = parse_user_date(str(data.get('date','')))
                paths = find_purchase_summaries(target)
                records = []
                seen = set()
                for path in paths:
                    workbook = load_workbook(path, read_only=True, data_only=True)
                    if 'Purchase Summary' not in workbook.sheetnames:
                        workbook.close()
                        continue
                    sheet = workbook['Purchase Summary']
                    headers = {str(sheet.cell(1,c).value or '').strip():c for c in range(1,sheet.max_column+1)}
                    invoice_col, date_col, case_col = headers.get('Invoice No.'), headers.get('Date'), headers.get('Case')
                    if not invoice_col or not date_col or not case_col:
                        workbook.close()
                        continue
                    totals = {}
                    for row in range(2, sheet.max_row+1):
                        if same_date(sheet.cell(row,date_col).value, target):
                            number = str(sheet.cell(row,invoice_col).value or '').strip()
                            if number:
                                totals[number] = totals.get(number,0.0) + float(sheet.cell(row,case_col).value or 0)
                    workbook.close()
                    for number, cases in totals.items():
                        if number not in seen:
                            seen.add(number)
                            records.append({'invoice':number, 'cases':cases})
                records.sort(key=lambda x:x['invoice'])
                return self.reply(200, {'date':target.strftime('%d-%m-%Y'), 'invoices':records, 'total_cases':sum(x['cases'] for x in records)})
            if self.path == '/api/settings':
                company = str(data.get('company','')).strip()
                if not company or len(company)>250:
                    raise ValueError('Select a company')
                SETTINGS.parent.mkdir(parents=True, exist_ok=True)
                SETTINGS.write_text(json.dumps({'company':company}), encoding='utf-8')
                return self.reply(200, {'ok':True})
            if self.path == '/api/jobs':
                if data.get('kind') not in {'mail','purchase','unloading','challans','stock','dms','connections'}:
                    raise ValueError('Unknown action')
                company = settings()['company']
                job = dict(id=uuid.uuid4().hex, created=datetime.now().isoformat(timespec='seconds'), status='queued',
                           company=company, request=data, detail='Waiting in queue', files=[])
                save_job(job)
                POOL.submit(run_job, job)
                return self.reply(202, job)
            self.reply(404, {'error':'Not found'})
        except Exception as exc:
            self.reply(400, {'error':str(exc)})

    def valid_host(self):
        host = self.headers.get('Host','').split(':')[0]
        return self.server.server_address[0] not in {'127.0.0.1','localhost'} or host in {'127.0.0.1','localhost'}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--host', default='127.0.0.1')
    parser.add_argument('--port', default=9001, type=int)
    a = parser.parse_args()
    if a.host not in {'127.0.0.1','localhost'} and not PASSWORD:
        parser.error('Set GOLDIEE_WEB_PASSWORD before enabling network access.')
    for job in jobs():
        if job['status'] in {'queued','running'}:
            job.update(status='failed', detail='Server restarted. Submit the operation again.')
            save_job(job)
    print(f'Goldiee dashboard: http://{a.host}:{a.port}', flush=True)
    ThreadingHTTPServer((a.host,a.port), Handler).serve_forever()


if __name__ == '__main__':
    main()
