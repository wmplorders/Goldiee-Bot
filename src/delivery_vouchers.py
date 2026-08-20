"""Excel-to-PDF delivery voucher generation for the Goldiee bot."""

from __future__ import annotations

import re
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from .delivery_voucher_template import (
    VoucherData,
    draw_voucher,
    prepare_stamp,
    register_fonts,
)
from .settings import PROJECT_ROOT


SHEET_NAME = "Sheet1"
HEADER_ROW = 10
STAMP_PATH = PROJECT_ROOT / "assets" / "delivery_voucher_stamp.png"

# Column B is intentionally used even though the workbook contains another
# Despatch Date column later in the sheet.
REQUIRED_COLUMNS = {
    "B": ("Despatch Date", "Date"),
    "C": ("Party Name & Station", "Invoice Party"),
    "D": ("Station", "Station"),
    "E": ("Nug", "Total Nug/Case"),
    "F": ("Bill No.", "Invoice No"),
    "H": ("Bill amount", "Total Invoice Value"),
    "I": ("Vch. No.", "Voucher No"),
    "N": ("Vehicle No.", "Vehicle No"),
    "P": ("Per case Rate", "Per Case Freight"),
    "Q": ("Freight Claimed", "Total Freight"),
}
OPTIONAL_COLUMNS = {
    "M": ("Eway No.", "Eway No"),
    "R": ("Remarks", "Remarks"),
}

ONES = [
    "Zero", "One", "Two", "Three", "Four", "Five", "Six", "Seven", "Eight", "Nine",
    "Ten", "Eleven", "Twelve", "Thirteen", "Fourteen", "Fifteen", "Sixteen",
    "Seventeen", "Eighteen", "Nineteen",
]
TENS = ["", "", "Twenty", "Thirty", "Forty", "Fifty", "Sixty", "Seventy", "Eighty", "Ninety"]


class DeliveryVoucherError(RuntimeError):
    """Raised when a workbook cannot be processed at all."""


def _under_hundred(number: int) -> str:
    if number < 20:
        return ONES[number]
    return TENS[number // 10] + (f"-{ONES[number % 10]}" if number % 10 else "")


def _under_thousand(number: int) -> str:
    words: list[str] = []
    if number >= 100:
        words.append(f"{ONES[number // 100]} Hundred")
        number %= 100
    if number:
        words.append(_under_hundred(number))
    return " ".join(words)


def integer_to_indian_words(number: int) -> str:
    if number == 0:
        return "Zero"
    if number < 0:
        return "Minus " + integer_to_indian_words(-number)
    words: list[str] = []
    for divisor, label in ((10_000_000, "Crore"), (100_000, "Lakh"), (1_000, "Thousand")):
        chunk, number = divmod(number, divisor)
        if chunk:
            words.append(f"{integer_to_indian_words(chunk)} {label}")
    if number:
        words.append(_under_thousand(number))
    return " ".join(words)


def amount_in_words(value: Decimal) -> str:
    amount = _decimal(value)
    rupees = int(amount)
    paise = int((amount - rupees) * 100)
    words = f"Rupees {integer_to_indian_words(rupees)}"
    if paise:
        words += f" and {integer_to_indian_words(paise)} Paise"
    return words + " Only"


def _decimal(value) -> Decimal:
    return Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _normalized_header(value) -> str:
    return " ".join(_text(value).casefold().split())


def _validate_structure(workbook) -> None:
    if SHEET_NAME not in workbook.sheetnames:
        raise DeliveryVoucherError(f"Worksheet `{SHEET_NAME}` was not found.")
    worksheet = workbook[SHEET_NAME]
    mismatches = []
    for column, (expected, _) in {**REQUIRED_COLUMNS, **OPTIONAL_COLUMNS}.items():
        actual = _text(worksheet[f"{column}{HEADER_ROW}"].value)
        if _normalized_header(actual) != _normalized_header(expected):
            mismatches.append(f"{column}{HEADER_ROW} should be `{expected}` (found `{actual or 'blank'}`)")
    if mismatches:
        raise DeliveryVoucherError("The Excel format is not recognized: " + "; ".join(mismatches))


def _read_rows(source: Path) -> tuple[list[VoucherData], list[str]]:
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        _validate_structure(workbook)
        worksheet = workbook[SHEET_NAME]
        vouchers: list[VoucherData] = []
        errors: list[str] = []
        for row in range(HEADER_ROW + 1, worksheet.max_row + 1):
            voucher_no = _text(worksheet[f"I{row}"].value)
            party = _text(worksheet[f"C{row}"].value)
            if not voucher_no and not party:
                continue

            raw = {
                "Date": worksheet[f"B{row}"].value,
                "Invoice Party": party,
                "Station": _text(worksheet[f"D{row}"].value),
                "Total Nug/Case": worksheet[f"E{row}"].value,
                "Invoice No": _text(worksheet[f"F{row}"].value),
                "Total Invoice Value": worksheet[f"H{row}"].value,
                "Voucher No": voucher_no,
                "Vehicle No": _text(worksheet[f"N{row}"].value),
                "Per Case Freight": worksheet[f"P{row}"].value,
                "Total Freight": worksheet[f"Q{row}"].value,
            }
            missing = [name for name, value in raw.items() if value in (None, "")]
            if missing:
                errors.append(f"Row {row}: missing {', '.join(missing)}")
                continue
            if not isinstance(raw["Date"], (date, datetime)):
                errors.append(f"Row {row}: invalid Despatch Date")
                continue

            try:
                invoice_value = _decimal(raw["Total Invoice Value"])
                vouchers.append(
                    VoucherData(
                        voucher_no=raw["Voucher No"],
                        voucher_date=raw["Date"],
                        invoice_no=raw["Invoice No"],
                        invoice_party=raw["Invoice Party"],
                        station=raw["Station"],
                        vehicle_no=raw["Vehicle No"],
                        eway_no=_text(worksheet[f"M{row}"].value),
                        total_cases=_decimal(raw["Total Nug/Case"]),
                        per_case_freight=_decimal(raw["Per Case Freight"]),
                        total_freight=_decimal(raw["Total Freight"]),
                        invoice_value=invoice_value,
                        invoice_value_words=amount_in_words(invoice_value),
                        remarks=_text(worksheet[f"R{row}"].value),
                    )
                )
            except Exception as exc:
                errors.append(f"Row {row}: invalid numeric value ({exc})")
        return vouchers, errors
    finally:
        workbook.close()


def default_output_filename(source: str | Path) -> str:
    stem = re.sub(r'[<>:"/\\|?*]+', "-", Path(source).stem).strip(" .")
    return f"{stem} Delivery Vouchers.pdf"


def build_delivery_vouchers(
    source: str | Path,
    output: str | Path,
    stamp_path: str | Path = STAMP_PATH,
) -> dict:
    source = Path(source)
    output = Path(output)
    stamp_path = Path(stamp_path)
    if source.suffix.casefold() not in {".xlsx", ".xlsm"}:
        raise DeliveryVoucherError("Delivery vouchers require an .xlsx or .xlsm file.")
    if not source.exists():
        raise DeliveryVoucherError(f"Excel file not found: {source}")
    if not stamp_path.exists():
        raise DeliveryVoucherError(f"Delivery voucher stamp not found: {stamp_path}")

    vouchers, errors = _read_rows(source)
    if not vouchers:
        detail = " " + " ".join(errors[:10]) if errors else ""
        raise DeliveryVoucherError("No valid voucher rows were found." + detail)

    output.parent.mkdir(parents=True, exist_ok=True)
    register_fonts()
    stamp = prepare_stamp(str(stamp_path))
    pdf = canvas.Canvas(str(output), pagesize=A4, pageCompression=1)
    pdf.setTitle(f"Delivery Vouchers - {source.stem}")
    pdf.setAuthor("Goenka Tradeventures")
    for voucher in vouchers:
        draw_voucher(pdf, voucher, stamp)
        pdf.showPage()
    pdf.save()
    return {
        "output": str(output),
        "generated": len(vouchers),
        "errors": errors,
        "source": str(source),
    }
