from __future__ import annotations

import argparse
import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from .settings import PROJECT_ROOT
from .tally_batch_stock import build as build_stock_status
from .tally_batch_stock import default_output_path as default_stock_output_path


SUMMARY_SEARCH_DIRS = [
    PROJECT_ROOT / "processed",
    PROJECT_ROOT / "downloads",
    Path.home() / "Downloads",
]
OUTPUT_DIR = PROJECT_ROOT / "processed"


def parse_user_date(value: str) -> datetime:
    value = value.strip()
    for fmt in ("%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d", "%d.%m.%Y"):
        try:
            return datetime.strptime(value, fmt)
        except ValueError:
            pass
    raise ValueError("Date must be like 30-05-2026 or 30/05/2026")


def normalize_item_for_stock(item: str, cases: float, pcs: float) -> str:
    name = str(item or "").strip()
    if name == "SEVIYAN PREMIUM 75g." and cases and round(pcs / cases) == 384:
        return "SEVIYAN PREMIUM 75g. - 384PKT"
    return name


def money_key(value: Any) -> str:
    if value in (None, ""):
        return ""
    try:
        return f"{float(value):.2f}"
    except (TypeError, ValueError):
        return str(value).strip()


def qty_text(case_qty: float, pcs_qty: float, pcs_per_case: float | None) -> str:
    if pcs_per_case and pcs_per_case > 0:
        whole_cases = int(pcs_qty // pcs_per_case)
        loose_pcs = round(pcs_qty - whole_cases * pcs_per_case)
        if loose_pcs:
            return f"{whole_cases} case, {loose_pcs:g} pcs"
        return f"{whole_cases} case"
    if float(case_qty).is_integer():
        return f"{int(case_qty)} case"
    return f"{case_qty:g} case"


def current_stock_text(stock_rows: list[dict[str, Any]]) -> str:
    pieces = []
    for row in sorted(stock_rows, key=lambda r: money_key(r["mrp"])):
        mrp = money_key(row["mrp"]) or "No MRP"
        pieces.append(f"Rs. {mrp} - {qty_text(row['case_qty'], row['qty'], row.get('pcs_per_case'))}")
    return ", ".join(pieces)


def find_purchase_summaries(target_date: datetime) -> list[Path]:
    candidates = []
    seen = set()
    for folder in SUMMARY_SEARCH_DIRS:
        if not folder.exists():
            continue
        for path in folder.glob("*.xlsx"):
            if path.name.startswith("~$"):
                continue
            resolved = str(path.resolve()).casefold()
            if resolved in seen:
                continue
            seen.add(resolved)
            if workbook_has_purchase_date(path, target_date):
                candidates.append(path)
    return sorted(candidates, key=lambda p: p.name.casefold())


def workbook_has_purchase_date(path: Path, target_date: datetime) -> bool:
    try:
        wb = load_workbook(path, data_only=True, read_only=True)
    except Exception:
        return False
    if "Purchase Summary" not in wb.sheetnames:
        return False
    ws = wb["Purchase Summary"]
    headers = {str(ws.cell(1, col).value or "").strip(): col for col in range(1, ws.max_column + 1)}
    date_col = headers.get("Date")
    if not date_col:
        return False
    for row in range(2, ws.max_row + 1):
        if same_date(ws.cell(row, date_col).value, target_date):
            return True
    return False


def same_date(value: Any, target_date: datetime) -> bool:
    if isinstance(value, datetime):
        return value.date() == target_date.date()
    text = str(value or "").strip()
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d.%m.%Y"):
        try:
            return datetime.strptime(text, fmt).date() == target_date.date()
        except ValueError:
            pass
    return False


def load_combined_purchase_rows(paths: list[Path], target_date: datetime) -> list[dict[str, Any]]:
    grouped: dict[tuple[str, str, str], dict[str, Any]] = {}
    for path in paths:
        wb = load_workbook(path, data_only=True, read_only=True)
        ws = wb["Purchase Summary"]
        headers = {str(ws.cell(1, col).value or "").strip(): col for col in range(1, ws.max_column + 1)}
        for row in range(2, ws.max_row + 1):
            if not same_date(ws.cell(row, headers["Date"]).value, target_date):
                continue
            item = str(ws.cell(row, headers["ITEM"]).value or "").strip()
            cases = float(ws.cell(row, headers["Case"]).value or 0)
            pcs = float(ws.cell(row, headers["QTY"]).value or 0)
            mapped_item = normalize_item_for_stock(item, cases, pcs)
            mrp = money_key(ws.cell(row, headers["MRP"]).value)
            ptr = money_key(ws.cell(row, headers["PTR"]).value)
            key = (mapped_item, mrp, ptr)
            if key not in grouped:
                grouped[key] = {
                    "item": mapped_item,
                    "purchase_item": item,
                    "mrp": mrp,
                    "ptr": ptr,
                    "cases": 0.0,
                    "pcs": 0.0,
                }
            grouped[key]["cases"] += cases
            grouped[key]["pcs"] += pcs
    return sorted(grouped.values(), key=lambda r: (r["item"].casefold(), r["mrp"], r["ptr"]))


def collect_invoice_numbers(paths: list[Path], target_date: datetime) -> list[str]:
    invoice_numbers = set()
    for path in paths:
        wb = load_workbook(path, data_only=True, read_only=True)
        ws = wb["Purchase Summary"]
        headers = {str(ws.cell(1, col).value or "").strip(): col for col in range(1, ws.max_column + 1)}
        date_col = headers.get("Date")
        invoice_col = headers.get("Invoice No.")
        if not date_col or not invoice_col:
            continue
        for row in range(2, ws.max_row + 1):
            if same_date(ws.cell(row, date_col).value, target_date):
                invoice_no = str(ws.cell(row, invoice_col).value or "").strip()
                if invoice_no:
                    invoice_numbers.add(invoice_no)
    return sorted(invoice_numbers)


def load_stock(stock_path: Path) -> dict[str, list[dict[str, Any]]]:
    wb = load_workbook(stock_path, data_only=True, read_only=True)
    ws = wb["Tally Batch Stock"]
    headers = {str(ws.cell(1, col).value or "").strip(): col for col in range(1, ws.max_column + 1)}
    by_item: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in range(2, ws.max_row + 1):
        item = str(ws.cell(row, headers["Stock Item"]).value or "").strip()
        if not item:
            continue
        by_item[item].append(
            {
                "item": item,
                "ptr": money_key(ws.cell(row, headers["Batch / PTR"]).value),
                "mrp": money_key(ws.cell(row, headers["MRP Rate"]).value),
                "qty": float(ws.cell(row, headers["Closing Qty"]).value or 0),
                "case_qty": float(ws.cell(row, headers["Closing Case Qty"]).value or 0),
                "pcs_per_case": float(ws.cell(row, headers["PCS per Case"]).value or 0) or None,
            }
        )
    return by_item


def classify_row(purchase_row: dict[str, Any], stock_by_item: dict[str, list[dict[str, Any]]]) -> tuple[str, str]:
    stock_rows = stock_by_item.get(purchase_row["item"], [])
    if not stock_rows:
        return "green", ""
    exact = [
        row
        for row in stock_rows
        if row["mrp"] == purchase_row["mrp"] and row["ptr"] == purchase_row["ptr"]
    ]
    stock_combos = {(row["mrp"], row["ptr"]) for row in stock_rows}
    stock_text = current_stock_text(stock_rows)
    if exact and len(stock_combos) == 1:
        return "normal", stock_text
    return "red", stock_text


def build_unloading_slip(target_date: datetime) -> dict[str, Any]:
    purchase_files = find_purchase_summaries(target_date)
    if not purchase_files:
        raise FileNotFoundError(
            f"No purchase summaries found for {target_date.strftime('%d-%m-%Y')} in processed/download folders."
        )
    purchase_rows = load_combined_purchase_rows(purchase_files, target_date)
    invoice_numbers = collect_invoice_numbers(purchase_files, target_date)
    stock_path = default_stock_output_path()
    build_stock_status(stock_path)
    stock_by_item = load_stock(stock_path)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    base = OUTPUT_DIR / f"Unloading Slip {target_date.strftime('%d-%m-%Y')}"
    xlsx_path = base.with_suffix(".xlsx")
    pdf_path = base.with_suffix(".pdf")

    wb = Workbook()
    ws = wb.active
    ws.title = "Unloading Slip"
    invoice_text = "Invoices: " + (", ".join(invoice_numbers) if invoice_numbers else "Not found")
    ws.append([invoice_text, "", "", "", "", "", "", ""])
    ws.append(["S. No.", "Item Name", "MRP", "Cases", "PCS", "", "", "Current Stock"])
    ws.merge_cells("A1:H1")
    ws["A1"].font = Font(bold=True)
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    green = PatternFill("solid", fgColor="C6EFCE")
    red = PatternFill("solid", fgColor="FFC7CE")
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[2]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    rows_for_pdf = []
    for serial, row in enumerate(purchase_rows, start=1):
        status, stock_text = classify_row(row, stock_by_item)
        values = [serial, row["item"], float(row["mrp"]), row["cases"], row["pcs"], "", "", stock_text]
        ws.append(values)
        excel_row = ws.max_row
        if status == "green":
            fill = green
        elif status == "red":
            fill = red
        else:
            fill = None
        if fill:
            for col in range(1, 9):
                ws.cell(excel_row, col).fill = fill
        rows_for_pdf.append((values, status))

    format_slip(ws)
    try:
        wb.save(xlsx_path)
    except PermissionError:
        timestamp = datetime.now().strftime("%H-%M-%S")
        xlsx_path = base.with_name(f"{base.name} {timestamp}").with_suffix(".xlsx")
        wb.save(xlsx_path)
    write_pdf(pdf_path, rows_for_pdf, target_date, invoice_numbers)
    return {
        "xlsx": str(xlsx_path),
        "pdf": str(pdf_path),
        "stock_status": str(stock_path),
        "purchase_files": [str(path) for path in purchase_files],
        "invoice_numbers": invoice_numbers,
        "rows": len(purchase_rows),
    }


def format_slip(ws) -> None:
    widths = {"A": 8, "B": 52, "C": 10, "D": 10, "E": 12, "F": 12, "G": 12, "H": 48}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=3):
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        row[0].alignment = Alignment(horizontal="center", vertical="center")
        row[2].number_format = "0.00"
        row[3].number_format = "0.00"
        row[4].number_format = "0"
    ws.freeze_panes = "A3"
    ws.sheet_view.showGridLines = False
    table = Table(displayName="UnloadingSlipTable", ref=f"A2:H{ws.max_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=False, showColumnStripes=False)
    ws.add_table(table)


def write_pdf(path: Path, rows: list[tuple[list[Any], str]], target_date: datetime, invoice_numbers: list[str]) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import landscape, A4
    from reportlab.pdfbase.pdfmetrics import stringWidth
    from reportlab.pdfgen import canvas

    page_width, page_height = landscape(A4)
    margin = 12
    title_height = 28
    table_top = page_height - margin - title_height - 4
    table_bottom = margin
    table_height = table_top - table_bottom
    row_count = len(rows) + 1
    row_height = table_height / max(row_count, 1)
    font_size = max(4.4, min(7.0, row_height * 0.42))
    header_font_size = max(5.0, min(7.5, row_height * 0.46))
    line_height = font_size + 1.1

    columns = [
        ("S. No.", 0.045),
        ("Item Name", 0.305),
        ("MRP", 0.065),
        ("Cases", 0.065),
        ("PCS", 0.075),
        ("", 0.065),
        ("", 0.065),
        ("Current Stock", 0.315),
    ]
    usable_width = page_width - 2 * margin
    col_widths = [usable_width * fraction for _, fraction in columns]

    pdf = canvas.Canvas(str(path), pagesize=landscape(A4))
    pdf.setTitle(f"Unloading Slip {target_date.strftime('%d-%m-%Y')}")

    def fit_line(text: str, width: float, size: float) -> str:
        text = str(text or "")
        if stringWidth(text, "Helvetica", size) <= width:
            return text
        ellipsis = "..."
        while text and stringWidth(text + ellipsis, "Helvetica", size) > width:
            text = text[:-1]
        return text + ellipsis if text else ellipsis

    pdf.setFont("Helvetica-Bold", 12)
    pdf.drawCentredString(page_width / 2, page_height - margin - 11, f"Unloading Slip {target_date.strftime('%d-%m-%Y')}")
    pdf.setFont("Helvetica", 7)
    invoice_text = "Invoices: " + (", ".join(invoice_numbers) if invoice_numbers else "Not found")
    pdf.drawString(margin, page_height - margin - 23, fit_line(invoice_text, page_width - 2 * margin, 7))

    def wrap_text(text: str, width: float, size: float, max_lines: int) -> list[str]:
        words = str(text or "").split()
        lines: list[str] = []
        current = ""
        for word in words:
            trial = f"{current} {word}".strip()
            if stringWidth(trial, "Helvetica", size) <= width:
                current = trial
            else:
                if current:
                    lines.append(current)
                current = word
                if len(lines) >= max_lines:
                    break
        if current and len(lines) < max_lines:
            lines.append(current)
        if not lines:
            lines = [""]
        if len(lines) == max_lines and words:
            consumed = " ".join(lines)
            if len(consumed) < len(str(text)):
                lines[-1] = fit_line(lines[-1], width, size)
        return lines[:max_lines]

    def draw_cell(x: float, y: float, width: float, height: float, text: Any, bold: bool = False, align: str = "left") -> None:
        pdf.rect(x, y, width, height, stroke=1, fill=0)
        size = header_font_size if bold else font_size
        font = "Helvetica-Bold" if bold else "Helvetica"
        pdf.setFont(font, size)
        max_lines = max(1, int((height - 3) // line_height))
        lines = wrap_text(str(text or ""), width - 4, size, max_lines)
        start_y = y + height - size - 2
        for idx, line in enumerate(lines):
            line_y = start_y - idx * line_height
            if align == "center":
                pdf.drawCentredString(x + width / 2, line_y, fit_line(line, width - 4, size))
            elif align == "right":
                pdf.drawRightString(x + width - 2, line_y, fit_line(line, width - 4, size))
            else:
                pdf.drawString(x + 2, line_y, fit_line(line, width - 4, size))

    y = table_top - row_height
    x = margin
    pdf.setFillColor(colors.HexColor("#1F4E78"))
    pdf.rect(margin, y, usable_width, row_height, stroke=0, fill=1)
    pdf.setFillColor(colors.white)
    for (header, _), width in zip(columns, col_widths):
        draw_cell(x, y, width, row_height, header, bold=True, align="center")
        x += width
    pdf.setFillColor(colors.black)

    for values, status in rows:
        y -= row_height
        if status == "green":
            pdf.setFillColor(colors.HexColor("#C6EFCE"))
            pdf.rect(margin, y, usable_width, row_height, stroke=0, fill=1)
        elif status == "red":
            pdf.setFillColor(colors.HexColor("#FFC7CE"))
            pdf.rect(margin, y, usable_width, row_height, stroke=0, fill=1)
        pdf.setFillColor(colors.black)
        rendered = [
            values[0],
            values[1],
            f"{values[2]:.2f}",
            f"{values[3]:g}",
            f"{values[4]:g}",
            "",
            "",
            values[7],
        ]
        x = margin
        for idx, (cell, width) in enumerate(zip(rendered, col_widths)):
            align = "center" if idx in {0, 2, 3, 4, 5, 6} else "left"
            draw_cell(x, y, width, row_height, cell, align=align)
            x += width

    pdf.showPage()
    pdf.save()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("date")
    args = parser.parse_args()
    result = build_unloading_slip(parse_user_date(args.date))
    print(result)


if __name__ == "__main__":
    main()
