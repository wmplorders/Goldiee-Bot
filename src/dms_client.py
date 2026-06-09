from __future__ import annotations

import json
import base64
from datetime import datetime
from pathlib import Path
from typing import Any

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from playwright.sync_api import sync_playwright

from .settings import PROJECT_ROOT


DMS_ORIGIN = "https://dms.fieldassist.io"
SESSION_PATH = PROJECT_ROOT / "data" / "dms_storage_state.json"
PROFILE_PATH = PROJECT_ROOT / "data" / "dms-browser-profile"


class DmsAuthenticationRequired(RuntimeError):
    pass


def authorize(config: dict) -> None:
    """One-time interactive authorization required by FieldAssist reCAPTCHA."""
    dms = config["dms"]
    PROFILE_PATH.mkdir(parents=True, exist_ok=True)
    with sync_playwright() as playwright:
        context = playwright.chromium.launch_persistent_context(
            str(PROFILE_PATH),
            executable_path=config["whatsapp"]["edge_executable"],
            headless=False,
        )
        page = context.pages[0] if context.pages else context.new_page()
        page.goto(f"{DMS_ORIGIN}/signin-oidc?company-id={dms['company_id']}")
        print("Complete DMS login in the opened Edge window.")
        page.wait_for_url(f"{DMS_ORIGIN}/**", timeout=300_000)
        page.goto(f"{DMS_ORIGIN}/signin-oidc?company-id={dms['company_id']}")
        page.wait_for_url(f"{DMS_ORIGIN}/**", timeout=120_000)
        _save_storage_state(context)
        context.close()


def fetch_inventory(config: dict) -> list[dict[str, Any]]:
    session = _authenticated_session(config)
    response = session.get(
        f"{DMS_ORIGIN}/api/v2/inventory/stocks",
        params={"category-info": "true"},
        timeout=120,
    )
    if response.status_code == 401:
        _renew_session_headlessly(config)
        session = _authenticated_session(config)
        response = session.get(
            f"{DMS_ORIGIN}/api/v2/inventory/stocks",
            params={"category-info": "true"},
            timeout=120,
        )
    response.raise_for_status()
    payload = response.json()
    if isinstance(payload, list):
        return payload
    for key in ("data", "stocks", "inventory", "items"):
        if isinstance(payload.get(key), list):
            return payload[key]
    raise RuntimeError("DMS inventory API returned an unsupported response format.")


def build_inventory_workbook(config: dict, output: Path | None = None) -> dict[str, Any]:
    rows = fetch_inventory(config)
    output = output or _default_output_path()
    output.parent.mkdir(parents=True, exist_ok=True)
    keys = _ordered_keys(rows)

    wb = Workbook()
    ws = wb.active
    ws.title = "DMS Current Stock"
    ws.append(keys)
    for row in rows:
        ws.append([_excel_value(row.get(key)) for key in keys])

    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    for column in ws.columns:
        letter = column[0].column_letter
        longest = max(len(str(cell.value or "")) for cell in column[:200])
        ws.column_dimensions[letter].width = min(max(longest + 2, 12), 45)
    if rows and keys:
        end = ws.cell(row=len(rows) + 1, column=len(keys)).coordinate
        table = Table(displayName="DmsCurrentStock", ref=f"A1:{end}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)
    wb.save(output)
    return {"output": str(output), "rows": len(rows), "columns": len(keys)}


def _authenticated_session(config: dict) -> requests.Session:
    if not SESSION_PATH.exists():
        raise DmsAuthenticationRequired(
            "DMS has not been authorized. Run `python -m src.main dms-authorize` once."
        )
    session = _session_from_storage()
    response = session.get(f"{DMS_ORIGIN}/api/v1/user/companies", timeout=30)
    if response.status_code == 401:
        _renew_session_headlessly(config)
        session = _session_from_storage()
        response = session.get(f"{DMS_ORIGIN}/api/v1/user/companies", timeout=30)
    response.raise_for_status()
    expected_id = int(config["dms"]["company_id"])
    company_ids = {
        int(company.get("companyId") or company.get("id"))
        for company in response.json()
        if company.get("companyId") or company.get("id")
    }
    if expected_id not in company_ids:
        raise RuntimeError(f"DMS company ID {expected_id} is not accessible to this account.")
    _switch_company(session, expected_id)
    return session


def _switch_company(session: requests.Session, company_id: int) -> None:
    access_token = session.cookies.get("accessToken", domain="dms.fieldassist.io")
    if access_token and _token_company_id(access_token) == company_id:
        session.headers["Authorization"] = f"Bearer {access_token}"
        return

    response = session.post(
        f"{DMS_ORIGIN}/api/v2/login/change-company/{company_id}",
        json={},
        timeout=60,
    )
    response.raise_for_status()
    tokens = response.json()
    access_token = tokens.get("accessToken")
    if not access_token or _token_company_id(access_token) != company_id:
        raise RuntimeError(f"DMS did not issue an access token for company ID {company_id}.")

    _store_token(session, "accessToken", access_token)
    if tokens.get("refreshToken"):
        _store_token(session, "refreshToken", tokens["refreshToken"])
    session.headers["Authorization"] = f"Bearer {access_token}"
    _persist_session_cookies(session)


def _token_company_id(token: str) -> int | None:
    try:
        payload = token.split(".")[1]
        payload += "=" * (-len(payload) % 4)
        claims = json.loads(base64.urlsafe_b64decode(payload))
        return int(claims["companyId"])
    except (IndexError, KeyError, TypeError, ValueError, json.JSONDecodeError):
        return None


def _store_token(session: requests.Session, name: str, value: str) -> None:
    session.cookies.set(name, value, domain="dms.fieldassist.io", path="/")


def _persist_session_cookies(session: requests.Session) -> None:
    state = json.loads(SESSION_PATH.read_text(encoding="utf-8"))
    cookies = {
        (cookie["domain"], cookie["path"], cookie["name"]): cookie
        for cookie in state.get("cookies", [])
    }
    for cookie in session.cookies:
        key = (cookie.domain, cookie.path, cookie.name)
        saved = cookies.get(key, {
            "name": cookie.name,
            "domain": cookie.domain,
            "path": cookie.path,
            "httpOnly": False,
            "secure": True,
            "sameSite": "Lax",
        })
        saved["value"] = cookie.value
        token_expiry = _token_expiry(cookie.value)
        if token_expiry:
            saved["expires"] = token_expiry
        cookies[key] = saved
    state["cookies"] = list(cookies.values())
    SESSION_PATH.write_text(json.dumps(state, indent=2), encoding="utf-8")


def _token_expiry(token: str) -> int | None:
    try:
        payload = token.split(".")[1]
        payload += "=" * (-len(payload) % 4)
        return int(json.loads(base64.urlsafe_b64decode(payload))["exp"])
    except (IndexError, KeyError, TypeError, ValueError, json.JSONDecodeError):
        return None


def _renew_session_headlessly(config: dict) -> None:
    dms = config["dms"]
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(
            executable_path=config["whatsapp"]["edge_executable"],
            headless=True,
        )
        context = browser.new_context(storage_state=str(SESSION_PATH))
        page = context.new_page()
        page.goto(
            f"{DMS_ORIGIN}/signin-oidc?company-id={dms['company_id']}",
            wait_until="domcontentloaded",
            timeout=120_000,
        )
        page.wait_for_timeout(5_000)
        if "login-v2.fieldassist.io" in page.url:
            browser.close()
            raise DmsAuthenticationRequired(
                "FieldAssist expired the identity session and requires reCAPTCHA. "
                "Run `python -m src.main dms-authorize` once to renew it."
            )
        _save_storage_state(context)
        browser.close()


def _save_storage_state(context) -> None:
    SESSION_PATH.parent.mkdir(parents=True, exist_ok=True)
    context.storage_state(path=str(SESSION_PATH))


def _session_from_storage() -> requests.Session:
    state = json.loads(SESSION_PATH.read_text(encoding="utf-8"))
    session = requests.Session()
    session.headers.update({"Accept": "application/json", "User-Agent": "Goldiee-Purchase-Bot/1.0"})
    for cookie in state.get("cookies", []):
        session.cookies.set(
            cookie["name"],
            cookie["value"],
            domain=cookie.get("domain"),
            path=cookie.get("path", "/"),
        )
    for origin in state.get("origins", []):
        if origin.get("origin") != DMS_ORIGIN:
            continue
        storage = {entry["name"]: entry["value"] for entry in origin.get("localStorage", [])}
        token = storage.get("accessToken")
        if token:
            session.headers["Authorization"] = f"Bearer {token}"
    return session


def _ordered_keys(rows: list[dict[str, Any]]) -> list[str]:
    preferred = [
        "productErpId", "productName", "variant", "productDivision", "primaryCategory",
        "secondaryCategory", "mrpBatchNo", "currentStockCases", "currentStockPieces",
        "currentStock", "mrp", "purchasePrice", "netValue", "warehouseErpId",
    ]
    found = {key for row in rows for key in row}
    return [key for key in preferred if key in found] + sorted(found.difference(preferred))


def _excel_value(value: Any) -> Any:
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return json.dumps(value, ensure_ascii=False, sort_keys=True)


def _default_output_path() -> Path:
    stamp = datetime.now().strftime("%Y-%m-%d %H-%M-%S")
    return PROJECT_ROOT / "downloads" / f"DMS Stock {stamp}.xlsx"
