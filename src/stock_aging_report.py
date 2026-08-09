from __future__ import annotations

import re
import xml.etree.ElementTree as ET
from copy import copy
from datetime import datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .settings import PROJECT_ROOT
from .tally_batch_stock import build as build_stock_status
from .tally_batch_stock import default_output_path as default_stock_output_path
from .tally_context import clean_tally_xml, current_company_static_variables
from .tally_context import get_tally_config, tally_request


DEFAULT_TEMPLATE = Path(r"C:\Users\goenk\Downloads\Stock Aging Report.xlsx")


def parse_report_date(value: str) -> datetime:
    value = value.strip()
    for fmt in ("%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d", "%Y/%m/%d"):
        try:
            return datetime.strptime(value, fmt)
        except ValueError:
            continue
    raise ValueError(f"Could not parse date {value!r}. Use DD-MM-YYYY or DD/MM/YYYY.")


def build_stock_aging_report(
    config: dict[str, Any],
    from_date: datetime,
    to_date: datetime,
    output: Path | None = None,
) -> dict[str, Any]:
    template = Path(config.get("stock_aging", {}).get("template") or DEFAULT_TEMPLATE)
    if not template.exists():
        raise FileNotFoundError(f"Stock ageing template not found: {template}")

    stock_result = build_stock_status(default_stock_output_path())
    stock_path = Path(stock_result["output"])
    item_meta = fetch_item_metadata()
    rows = aggregate_stock_rows(stock_path, item_meta)

    output = output or default_output_path(config, from_date, to_date)
    output.parent.mkdir(parents=True, exist_ok=True)
    fill_template(template, output, rows, from_date, to_date)

    report_items = {row["item"] for row in rows}
    missing_codes = sorted(row["item"] for row in rows if not row["code"])
    missing_weights = sorted(row["item"] for row in rows if row["weight_gm"] is None)
    return {
        "output": str(output),
        "stock_source": str(stock_path),
        "rows": len(rows),
        "tally_positive_items": len(report_items),
        "missing_codes": missing_codes,
        "missing_weights": missing_weights,
        "from_date": from_date.strftime("%d-%m-%Y"),
        "to_date": to_date.strftime("%d-%m-%Y"),
    }


def default_output_path(config: dict[str, Any], from_date: datetime, to_date: datetime) -> Path:
    processed = Path(config["paths"]["processed"])
    return processed / (
        "Stock Aging Report Goldiee "
        f"{from_date.strftime('%d-%m-%Y')} to {to_date.strftime('%d-%m-%Y')}.xlsx"
    )


def fetch_item_metadata() -> dict[str, dict[str, Any]]:
    tally_config = get_tally_config()
    request_xml = (
        "<ENVELOPE><HEADER><VERSION>1</VERSION><TALLYREQUEST>Export</TALLYREQUEST>"
        "<TYPE>Collection</TYPE><ID>Goldiee Bot Stock Items For Ageing</ID></HEADER><BODY><DESC>"
        "<STATICVARIABLES><SVEXPORTFORMAT>$$SysName:XML</SVEXPORTFORMAT>"
        f"{current_company_static_variables(tally_config)}</STATICVARIABLES>"
        "<TDL><TDLMESSAGE><COLLECTION NAME=\"Goldiee Bot Stock Items For Ageing\">"
        "<TYPE>Stock Item</TYPE>"
        "<FETCH>Name,Parent,BaseUnits,LANGUAGENAME.LIST.*,NAME.LIST.*</FETCH>"
        "</COLLECTION></TDLMESSAGE></TDL></DESC></BODY></ENVELOPE>"
    )
    root = ET.fromstring(clean_tally_xml(tally_request(request_xml, tally_config, timeout=60)))
    metadata: dict[str, dict[str, Any]] = {}
    for item in root.findall(".//STOCKITEM"):
        name = item.attrib.get("NAME", "").strip()
        if not name:
            continue
        aliases = [
            elem.text.strip()
            for elem in item.findall(".//NAME")
            if elem.text and elem.text.strip() and elem.text.strip() != name
        ]
        metadata[name] = {
            "code": aliases[0] if aliases else "",
            "category": text(item, "PARENT"),
            "base_unit": text(item, "BASEUNITS"),
        }
    return metadata


def aggregate_stock_rows(stock_path: Path, item_meta: dict[str, dict[str, Any]]) -> list[dict[str, Any]]:
    workbook = load_workbook(stock_path, data_only=True, read_only=True)
    sheet = workbook["Tally Batch Stock"]
    rows_iter = sheet.iter_rows(values_only=True)
    headers = list(next(rows_iter))
    index = {header: pos for pos, header in enumerate(headers)}
    aggregated: dict[str, dict[str, Any]] = {}
    for values in rows_iter:
        item = values[index["Stock Item"]]
        qty = abs_decimal(values[index["Closing Qty"]])
        if not item or qty <= 0:
            continue
        item = str(item)
        current = aggregated.setdefault(
            item,
            {
                "item": item,
                "pcs": Decimal("0"),
                "cases": Decimal("0"),
                "value": Decimal("0"),
                "rate_qty": Decimal("0"),
            },
        )
        current["pcs"] += qty
        current["cases"] += abs_decimal(values[index["Closing Case Qty"]])
        current["value"] += abs_decimal(values[index["Closing Value"]])
        rate = abs_decimal(values[index["Closing Rate"]])
        if rate:
            current["rate_qty"] += rate * qty
    workbook.close()

    result = []
    for item, row in sorted(aggregated.items(), key=lambda pair: pair[0].casefold()):
        meta = item_meta.get(item, {})
        qty = row["pcs"]
        value = row["value"]
        rate = value / qty if qty and value else (row["rate_qty"] / qty if qty and row["rate_qty"] else Decimal("0"))
        result.append(
            {
                "code": meta.get("code", ""),
                "item": item,
                "category": meta.get("category", ""),
                "weight_gm": parse_weight_gm(item),
                "pcs": qty,
                "cases": row["cases"],
                "rate": rate,
            }
        )
    return result


def fill_template(
    template: Path,
    output: Path,
    rows: list[dict[str, Any]],
    from_date: datetime,
    to_date: datetime,
) -> None:
    workbook = load_workbook(template)
    sheet = workbook.active
    if sheet.max_row >= 6:
        sheet.delete_rows(6, sheet.max_row - 5)

    style_workbook = load_workbook(template)
    style_sheet = style_workbook.active
    style_cells = [style_sheet.cell(6, column) for column in range(1, 17)]

    for row_no, row in enumerate(rows, start=6):
        copy_row_style(sheet, row_no, style_cells)
        sheet.cell(row_no, 1).value = row["code"]
        sheet.cell(row_no, 2).value = row["item"]
        sheet.cell(row_no, 3).value = row["category"]
        sheet.cell(row_no, 4).value = number(row["weight_gm"])
        sheet.cell(row_no, 5).value = number(row["pcs"])
        sheet.cell(row_no, 7).value = number(row["cases"])
        sheet.cell(row_no, 9).value = number(row["rate"])
        sheet.cell(row_no, 10).value = "PC"
        sheet.cell(row_no, 11).value = f"=E{row_no}*I{row_no}"

    sheet["A4"] = (
        "STOCK DATE  (AS ON) : "
        f"{to_date.strftime('%d-%m-%Y')} "
        f"(Period: {from_date.strftime('%d-%m-%Y')} to {to_date.strftime('%d-%m-%Y')})"
    )
    sheet.freeze_panes = "A6"
    workbook.save(output)


def copy_row_style(sheet, row_no: int, style_cells) -> None:
    for column in range(1, 17):
        source = style_cells[column - 1]
        target = sheet.cell(row_no, column)
        if source.has_style:
            target._style = copy(source._style)
        target.number_format = source.number_format
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.border = copy(source.border)
        target.alignment = copy(source.alignment)


def text(node, tag: str) -> str:
    child = node.find(tag)
    return (child.text or "").strip() if child is not None and child.text else ""


def abs_decimal(value) -> Decimal:
    return abs(decimal_value(value))


def decimal_value(value) -> Decimal:
    if value in (None, ""):
        return Decimal("0")
    try:
        return Decimal(str(value).replace(",", "").strip())
    except InvalidOperation:
        return Decimal("0")


def number(value):
    if value in (None, ""):
        return None
    value = Decimal(value).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    return int(value) if value == value.to_integral() else float(value)


def parse_weight_gm(name: str):
    clean = re.sub(r"\[[^\]]*]|\([^)]*pcs[^)]*\)", " ", name, flags=re.IGNORECASE)
    total = Decimal("0")
    found = False
    for match in re.finditer(
        r"(?<![A-Z0-9])([0-9]+(?:\.[0-9]+)?)\s*(KG|KGS|KILO|GM|GMS|G|GRAM|GRAMS)\b",
        clean,
        flags=re.IGNORECASE,
    ):
        value = Decimal(match.group(1))
        unit = match.group(2).upper()
        if unit in {"KG", "KGS", "KILO"}:
            value *= Decimal("1000")
        total += value
        found = True
    return total if found else None
