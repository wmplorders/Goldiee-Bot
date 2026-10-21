from __future__ import annotations

import csv
import html
import re
import xml.etree.ElementTree as ET
from io import StringIO
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from .tally_context import current_company_static_variables, ensure_configured_company
from .tally_context import get_tally_config, tally_request
from .tally_job_lock import serialized


ROOT = Path(__file__).resolve().parents[1]
DOWNLOADS = ROOT / "downloads"
STOCK_XML = DOWNLOADS / "tally_stock_items_conversion.xml"
VOUCHER_XML = DOWNLOADS / "tally_inventory_vouchers_batches.xml"


def clean_xml(path: Path, ns: bool = False) -> str:
    text = path.read_text(encoding="utf-8", errors="ignore")
    text = re.sub(r"&#(?:[0-8]|1[0-9]|2[0-9]|3[0-1]);", "", text)
    if ns:
        text = text.replace("<ENVELOPE>", '<ENVELOPE xmlns:UDF="TallyUDF">', 1)
    return text


def text(node, tag: str) -> str:
    child = node.find(tag)
    return (child.text or "").strip() if child is not None else ""


def number(value):
    if value in (None, ""):
        return ""
    try:
        return float(str(value).replace(",", "").split("/")[0].strip())
    except ValueError:
        return value


def batch_sort_key(value):
    value = str(value or "").strip()
    if re.fullmatch(r"[-+]?\d+(?:\.\d+)?", value):
        return (0, float(value))
    return (1, value.casefold())


def parse_qty(value: str) -> tuple[float, str]:
    match = re.match(r"\s*([-+]?[0-9,.]+)\s+([^=]+)", value or "")
    if not match:
        return 0.0, ""
    return float(match.group(1).replace(",", "")), match.group(2).strip()


def udf_value(node, local_name: str):
    for elem in node.iter():
        if elem.tag.endswith(local_name):
            value = (elem.text or "").strip()
            if value:
                return number(value)
    return ""


def fetch_item_batch_summary(item: str) -> list[list[str]]:
    tally_config = get_tally_config()
    body = (
        "<ENVELOPE><HEADER><VERSION>1</VERSION><TALLYREQUEST>Export</TALLYREQUEST>"
        "<TYPE>Data</TYPE><ID>Item Batch Summary</ID></HEADER><BODY><DESC>"
        "<STATICVARIABLES><SVEXPORTFORMAT>$$SysName:ASCII</SVEXPORTFORMAT>"
        f"{current_company_static_variables(tally_config)}"
        f"<SVStockItemName>{html.escape(item)}</SVStockItemName>"
        f"<StockItemName>{html.escape(item)}</StockItemName>"
        "<EXPLODEFLAG>Yes</EXPLODEFLAG></STATICVARIABLES></DESC></BODY></ENVELOPE>"
    )
    data = tally_request(body, tally_config, timeout=25)
    return list(csv.reader(StringIO(data)))


def refresh_source_xml() -> None:
    tally_config = get_tally_config()
    DOWNLOADS.mkdir(parents=True, exist_ok=True)
    stock_body = (
        "<ENVELOPE><HEADER><VERSION>1</VERSION><TALLYREQUEST>Export</TALLYREQUEST>"
        "<TYPE>Collection</TYPE><ID>Goldiee Bot Stock Items Conversion</ID></HEADER><BODY><DESC>"
        "<STATICVARIABLES><SVEXPORTFORMAT>$$SysName:XML</SVEXPORTFORMAT>"
        f"{current_company_static_variables(tally_config)}</STATICVARIABLES>"
        "<TDL><TDLMESSAGE><COLLECTION NAME=\"Goldiee Bot Stock Items Conversion\">"
        "<TYPE>Stock Item</TYPE>"
        "<FETCH>Name,Parent,BaseUnits,AdditionalUnits,Denominator,Conversion,ClosingBalance</FETCH>"
        "</COLLECTION></TDLMESSAGE></TDL></DESC></BODY></ENVELOPE>"
    )
    STOCK_XML.write_text(tally_request(stock_body, tally_config, timeout=60), encoding="utf-8")

    voucher_body = (
        "<ENVELOPE><HEADER><VERSION>1</VERSION><TALLYREQUEST>Export</TALLYREQUEST>"
        "<TYPE>Collection</TYPE><ID>Goldiee Bot Inventory Voucher Batches</ID></HEADER><BODY><DESC>"
        "<STATICVARIABLES><SVEXPORTFORMAT>$$SysName:XML</SVEXPORTFORMAT>"
        f"{current_company_static_variables(tally_config)}</STATICVARIABLES>"
        "<TDL><TDLMESSAGE><COLLECTION NAME=\"Goldiee Bot Inventory Voucher Batches\">"
        "<TYPE>Voucher</TYPE>"
        "<FETCH>Date,IsCancelled,IsOptional,AllInventoryEntries.*,BatchAllocations.*</FETCH>"
        "</COLLECTION></TDLMESSAGE></TDL></DESC></BODY></ENVELOPE>"
    )
    VOUCHER_XML.write_text(tally_request(voucher_body, tally_config, timeout=120), encoding="utf-8")


@serialized
def build(output: Path) -> dict:
    ensure_configured_company()
    refresh_source_xml()
    stock_root = ET.fromstring(clean_xml(STOCK_XML))
    items = []
    item_meta = {}
    for item_node in stock_root.findall(".//STOCKITEM"):
        name = item_node.attrib.get("NAME", "").strip()
        if not name:
            continue
        closing_raw = text(item_node, "CLOSINGBALANCE")
        denominator = number(text(item_node, "DENOMINATOR"))
        conversion = number(text(item_node, "CONVERSION"))
        pcs_per_case = ""
        if isinstance(denominator, float) and isinstance(conversion, float) and conversion:
            pcs_per_case = denominator / conversion
        if closing_raw and not closing_raw.startswith("0.00"):
            items.append(name)
        item_meta[name] = {
            "unit": text(item_node, "BASEUNITS"),
            "case_unit": text(item_node, "ADDITIONALUNITS"),
            "denominator": denominator,
            "conversion": conversion,
            "pcs_per_case": pcs_per_case,
        }

    udf_map = {}
    voucher_root = ET.fromstring(clean_xml(VOUCHER_XML, ns=True))
    for voucher in voucher_root.findall(".//VOUCHER"):
        date = text(voucher, "DATE")
        for inv in voucher.findall("ALLINVENTORYENTRIES.LIST"):
            item = text(inv, "STOCKITEMNAME")
            if not item:
                continue
            for batch in inv.findall("BATCHALLOCATIONS.LIST"):
                batch_name = text(batch, "BATCHNAME")
                if not batch_name:
                    continue
                mrp = udf_value(batch, "UDF_AA_BATCH_MRP_RATE")
                sell = udf_value(batch, "UDF_AA_BATCH_SELL_RATE")
                if mrp != "" or sell != "":
                    old = udf_map.get((item, batch_name))
                    if old is None or date >= old["date"]:
                        udf_map[(item, batch_name)] = {"date": date, "mrp": mrp, "sell": sell}

    rows = []
    errors = []
    for item in items:
        try:
            batch_rows = fetch_item_batch_summary(item)
        except Exception as exc:
            errors.append([item, str(exc)])
            continue
        for rec in batch_rows:
            if len(rec) < 6:
                continue
            batch_name = rec[0].strip()
            if not batch_name or batch_name.casefold() == "grand total":
                continue
            qty, unit = parse_qty(rec[3].strip())
            meta = item_meta[item]
            pcs_per_case = meta["pcs_per_case"]
            case_qty = round(qty / pcs_per_case, 4) if pcs_per_case else ""
            udf = udf_map.get((item, batch_name), {})
            rows.append(
                [
                    item,
                    batch_name,
                    udf.get("mrp", ""),
                    udf.get("sell", ""),
                    qty,
                    unit or meta["unit"],
                    case_qty,
                    meta["case_unit"],
                    pcs_per_case,
                    meta["denominator"],
                    meta["conversion"],
                    number(rec[4].strip()),
                    number(rec[5].strip()),
                    udf.get("date", ""),
                ]
            )

    rows.sort(
        key=lambda row: (
            str(row[0]).casefold(),
            batch_sort_key(row[1]),
        )
    )
    write_workbook(output, rows, errors, len(items))
    return {
        "output": str(output),
        "items": len(items),
        "rows": len(rows),
        "errors": len(errors),
        "mrp_rows": sum(1 for row in rows if row[2] != ""),
        "sell_rows": sum(1 for row in rows if row[3] != ""),
    }


def write_workbook(output: Path, rows, errors, item_count: int) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Tally Batch Stock"
    ws.append(
        [
            "Stock Item",
            "Batch / PTR",
            "MRP Rate",
            "Sell Rate",
            "Closing Qty",
            "Unit",
            "Closing Case Qty",
            "Case Unit",
            "PCS per Case",
            "Denominator",
            "Conversion",
            "Closing Rate",
            "Closing Value",
            "MRP/Sell Source Date",
        ]
    )
    for row in rows:
        ws.append(row)

    summary = wb.create_sheet("Summary")
    summary.append(["Metric", "Value"])
    summary.append(["Source report", "Item Batch Summary"])
    summary.append(["Items queried", item_count])
    summary.append(["Batch rows returned", len(rows)])
    summary.append(["Rows with MRP from UDF", sum(1 for row in rows if row[2] != "")])
    summary.append(["Rows with Sell Rate from UDF", sum(1 for row in rows if row[3] != "")])
    summary.append(["PCS per Case source", "Stock item master DENOMINATOR / CONVERSION"])
    summary.append(["Errors", len(errors)])

    if errors:
        error_sheet = wb.create_sheet("Errors")
        error_sheet.append(["Stock Item", "Error"])
        for error in errors:
            error_sheet.append(error)

    style(ws, {"A": 52, "B": 12, "C": 12, "D": 12, "E": 14, "F": 10, "G": 16, "H": 10, "I": 12, "J": 12, "K": 12, "L": 14, "M": 15, "N": 16})
    style(summary, {"A": 32, "B": 70})
    if errors:
        style(wb["Errors"], {"A": 52, "B": 80})
    wb.save(output)


def style(sheet, widths: dict[str, int]) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="center")
            if isinstance(cell.value, float):
                cell.number_format = "0.00"
    if sheet.max_row > 1:
        n = sheet.max_column
        letters = ""
        while n:
            n, remainder = divmod(n - 1, 26)
            letters = chr(65 + remainder) + letters
        table_name = "Table" + re.sub(r"\W+", "", sheet.title)
        table = Table(displayName=table_name, ref=f"A1:{letters}{sheet.max_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
        sheet.add_table(table)


def default_output_path() -> Path:
    stamp = datetime.now().strftime("%Y-%m-%d %H-%M-%S")
    return DOWNLOADS / f"Stock Status {stamp}.xlsx"


if __name__ == "__main__":
    output_path = default_output_path()
    print(build(output_path))
