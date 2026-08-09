from __future__ import annotations

import json
import math
import re
import xml.etree.ElementTree as ET
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
DOWNLOADS = ROOT / "downloads"
AS_OF = datetime(2026, 7, 21)
STOCK_FILE = DOWNLOADS / "Stock Status 2026-07-21 11-53-49.xlsx"
VOUCHER_FILE = DOWNLOADS / "tally_inventory_vouchers_batches_20250401_20260721.xml"
OUTPUT_FILE = DOWNLOADS / "Tally Stock Aging Analysis 2026-07-21.xlsx"

BUCKETS = [
    ("180+ Days", 180, math.inf),
    ("150+ Days", 150, 180),
    ("120+ Days", 120, 150),
    ("90+ Days", 90, 120),
    ("60+ Days", 60, 90),
    ("30+ Days", 30, 60),
    ("0-30 Days", 0, 30),
]

PURCHASE_TYPES = {"Purchase", "GST Credit Note"}
SALE_TYPES = {"Tax Invoice", "Detergent Invoice", "Bill of Supply"}


def clean_xml(text: str) -> str:
    text = re.sub(r"&#(?:[0-8]|1[0-9]|2[0-9]|3[0-1]);", "", text)
    return text.replace("<ENVELOPE>", '<ENVELOPE xmlns:UDF="TallyUDF">', 1)


def text(node, tag: str) -> str:
    child = node.find(tag)
    return (child.text or "").strip() if child is not None and child.text else ""


def parse_date(value: str):
    if not re.fullmatch(r"\d{8}", value or ""):
        return None
    return datetime.strptime(value, "%Y%m%d")


def parse_qty(value: str) -> float:
    match = re.match(r"\s*([-+]?[0-9,.]+)", value or "")
    return float(match.group(1).replace(",", "")) if match else 0.0


def norm_key(item: str, ptr: str) -> tuple[str, str]:
    return (str(item or "").strip(), str(ptr or "").strip())


def read_current_stock():
    wb = load_workbook(STOCK_FILE, read_only=True, data_only=True)
    ws = wb["Tally Batch Stock"]
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    index = {name: i for i, name in enumerate(headers)}
    rows = []
    for rec in ws.iter_rows(min_row=2, values_only=True):
        item = rec[index["Stock Item"]]
        ptr = rec[index["Batch / PTR"]]
        qty = float(rec[index["Closing Qty"]] or 0)
        if not item or qty <= 0:
            continue
        rows.append(
            {
                "item": str(item).strip(),
                "ptr": str(ptr).strip(),
                "qty_pcs": qty,
                "qty_cases": float(rec[index["Closing Case Qty"]] or 0),
                "pcs_per_case": float(rec[index["PCS per Case"]] or 0),
            }
        )
    return rows


def read_vouchers():
    root = ET.fromstring(clean_xml(VOUCHER_FILE.read_text(encoding="utf-8", errors="ignore")))
    purchases = defaultdict(list)
    last_sale = {}
    date_values = []
    voucher_counts = defaultdict(int)
    for voucher in root.findall(".//VOUCHER"):
        date = parse_date(text(voucher, "DATE"))
        if not date:
            continue
        date_values.append(date)
        vtype = text(voucher, "VOUCHERTYPENAME")
        voucher_counts[vtype] += 1
        vno = text(voucher, "VOUCHERNUMBER")
        for inv in voucher.findall("ALLINVENTORYENTRIES.LIST"):
            item = text(inv, "STOCKITEMNAME")
            if not item:
                continue
            for batch in inv.findall("BATCHALLOCATIONS.LIST"):
                ptr = text(batch, "BATCHNAME")
                if not ptr:
                    continue
                qty = parse_qty(text(batch, "BILLEDQTY") or text(batch, "ACTUALQTY"))
                if qty <= 0:
                    continue
                key = norm_key(item, ptr)
                if vtype in PURCHASE_TYPES:
                    purchases[key].append({"date": date, "qty_pcs": qty, "voucher": vno})
                elif vtype in SALE_TYPES and (key not in last_sale or date > last_sale[key]):
                    last_sale[key] = date
    return purchases, last_sale, date_values, voucher_counts


def bucket_for_age(age_days: int) -> str:
    for name, low, high in BUCKETS:
        if age_days >= low and age_days < high:
            return name
    return "0-30 Days"


def case_qty(qty_pcs: float, pcs_per_case: float) -> float:
    return round(qty_pcs / pcs_per_case, 4) if pcs_per_case else 0


def remark_for(bucket: str, unmatched: bool = False) -> str:
    if unmatched:
        return "Review: opening/unmatched stock"
    if bucket == "180+ Days":
        return "Very old / slow"
    if bucket in {"150+ Days", "120+ Days"}:
        return "Old stock"
    if bucket == "90+ Days":
        return "Slow / review"
    if bucket == "60+ Days":
        return "Watch"
    if bucket == "30+ Days":
        return "Normal"
    return "Fresh"


def allocate_fifo(current_rows, purchases, last_sale):
    all_allocations = []
    unmatched_rows = []
    for stock in current_rows:
        key = norm_key(stock["item"], stock["ptr"])
        remaining = stock["qty_pcs"]
        stock_allocations = []
        lots = sorted(purchases.get(key, []), key=lambda lot: lot["date"], reverse=True)
        for lot in lots:
            if remaining <= 0:
                break
            take = min(remaining, lot["qty_pcs"])
            age = (AS_OF - lot["date"]).days
            bucket = bucket_for_age(age)
            row = {
                "bucket": bucket,
                "item": stock["item"],
                "ptr": stock["ptr"],
                "qty_pcs": round(take, 3),
                "qty_cases": case_qty(take, stock["pcs_per_case"]),
                "purchase_date": lot["date"],
                "age_days": age,
                "last_sale_date": last_sale.get(key),
                "remark": remark_for(bucket),
                "source": lot["voucher"],
            }
            all_allocations.append(row)
            stock_allocations.append(row)
            remaining -= take
        if remaining > 0.0001:
            row = {
                "bucket": "Opening / Unmatched",
                "item": stock["item"],
                "ptr": stock["ptr"],
                "qty_pcs": round(remaining, 3),
                "qty_cases": case_qty(remaining, stock["pcs_per_case"]),
                "purchase_date": None,
                "age_days": None,
                "last_sale_date": last_sale.get(key),
                "remark": remark_for("Opening / Unmatched", unmatched=True),
                "source": "Not matched in exported purchase history",
            }
            all_allocations.append(row)
            stock_allocations.append(row)
            unmatched_rows.append(row)
        stock["allocations"] = stock_allocations
    return all_allocations, unmatched_rows


def fmt_date(value):
    return value.strftime("%Y-%m-%d") if value else ""


def write_rows(ws, headers, rows):
    ws.append(headers)
    for row in rows:
        ws.append(row)
    if rows:
        end_col = ws.cell(row=1, column=len(headers)).column_letter
        table_name = re.sub(r"[^A-Za-z0-9_]", "", ws.title)[:20] + "Table"
        table = Table(displayName=table_name, ref=f"A1:{end_col}{len(rows) + 1}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
        ws.add_table(table)


def style_sheet(ws, widths):
    header_fill = PatternFill("solid", fgColor="1F4E78")
    thin = Side(style="thin", color="D9E2F3")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(vertical="top", wrap_text=False)
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.35
    ws.page_margins.bottom = 0.35


def build_workbook(current_rows, allocations, unmatched, dates, counts, last_sale):
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"

    totals = defaultdict(lambda: {"rows": 0, "pcs": 0.0, "cases": 0.0})
    for row in allocations:
        totals[row["bucket"]]["rows"] += 1
        totals[row["bucket"]]["pcs"] += row["qty_pcs"]
        totals[row["bucket"]]["cases"] += row["qty_cases"]

    summary_rows = [
        ["Metric", "Value"],
        ["As of date", AS_OF.strftime("%Y-%m-%d")],
        ["Current stock source", str(STOCK_FILE)],
        ["Voucher movement source", str(VOUCHER_FILE)],
        ["Voucher date range available", f"{min(dates).strftime('%Y-%m-%d')} to {max(dates).strftime('%Y-%m-%d')}" if dates else ""],
        ["Current stock rows", len(current_rows)],
        ["Total current PCS", round(sum(r["qty_pcs"] for r in current_rows), 3)],
        ["Total current Cases", round(sum(r["qty_cases"] for r in current_rows), 4)],
        ["FIFO note", "Remaining stock allocated against latest purchase lots first; older lots are assumed sold first."],
        ["Unmatched note", "Opening / Unmatched rows could not be linked to purchases in exported Tally movement."],
    ]
    for row in summary_rows:
        summary.append(row)
    summary.append([])
    summary.append(["Bucket", "Rows", "Qty PCS", "Qty Cases"])
    for bucket, _, _ in BUCKETS:
        t = totals[bucket]
        summary.append([bucket, t["rows"], round(t["pcs"], 3), round(t["cases"], 4)])
    t = totals["Opening / Unmatched"]
    summary.append(["Opening / Unmatched", t["rows"], round(t["pcs"], 3), round(t["cases"], 4)])
    summary.append([])
    summary.append(["Voucher Type", "Count"])
    for k, v in sorted(counts.items()):
        summary.append([k or "(blank)", v])
    style_sheet(summary, {"A": 32, "B": 120, "C": 16, "D": 16})

    headers = ["Item Name", "PTR / Batch", "Qty PCS", "Qty Cases", "Purchase Date", "Age Days", "Last Sale Date", "Remark", "Source Voucher"]
    for bucket, _, _ in BUCKETS:
        ws = wb.create_sheet(bucket)
        rows = [
            [
                r["item"],
                r["ptr"],
                r["qty_pcs"],
                r["qty_cases"],
                fmt_date(r["purchase_date"]),
                r["age_days"] if r["age_days"] is not None else "",
                fmt_date(r["last_sale_date"]),
                r["remark"],
                r["source"],
            ]
            for r in sorted((x for x in allocations if x["bucket"] == bucket), key=lambda x: (x["item"].casefold(), x["ptr"]))
        ]
        write_rows(ws, headers, rows)
        style_sheet(ws, {"A": 48, "B": 12, "C": 12, "D": 12, "E": 14, "F": 10, "G": 14, "H": 18, "I": 18})

    ws = wb.create_sheet("Opening Unmatched")
    rows = [
        [r["item"], r["ptr"], r["qty_pcs"], r["qty_cases"], "", "", fmt_date(r["last_sale_date"]), r["remark"], r["source"]]
        for r in sorted(unmatched, key=lambda x: (x["item"].casefold(), x["ptr"]))
    ]
    write_rows(ws, headers, rows)
    style_sheet(ws, {"A": 48, "B": 12, "C": 12, "D": 12, "E": 14, "F": 10, "G": 14, "H": 30, "I": 38})

    all_ws = wb.create_sheet("All Current Stock")
    all_headers = ["Item Name", "PTR / Batch", "Current Qty PCS", "Current Qty Cases", "PCS per Case", "Oldest Allocated Date", "Latest Allocated Date", "Main Bucket", "Last Sale Date", "Remark"]
    all_rows = []
    for stock in sorted(current_rows, key=lambda x: (x["item"].casefold(), x["ptr"])):
        allocs = stock.get("allocations", [])
        dated = [r for r in allocs if r["purchase_date"]]
        by_qty = defaultdict(float)
        remarks = []
        for row in allocs:
            by_qty[row["bucket"]] += row["qty_pcs"]
            if row["remark"] not in remarks:
                remarks.append(row["remark"])
        main_bucket = max(by_qty.items(), key=lambda kv: kv[1])[0] if by_qty else ""
        key = norm_key(stock["item"], stock["ptr"])
        all_rows.append(
            [
                stock["item"],
                stock["ptr"],
                stock["qty_pcs"],
                stock["qty_cases"],
                stock["pcs_per_case"],
                fmt_date(min((r["purchase_date"] for r in dated), default=None)),
                fmt_date(max((r["purchase_date"] for r in dated), default=None)),
                main_bucket,
                fmt_date(last_sale.get(key)),
                "; ".join(remarks),
            ]
        )
    write_rows(all_ws, all_headers, all_rows)
    style_sheet(all_ws, {"A": 48, "B": 12, "C": 14, "D": 14, "E": 12, "F": 16, "G": 16, "H": 18, "I": 14, "J": 34})

    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = "#,##0.00"
                elif isinstance(cell.value, int):
                    cell.number_format = "#,##0"
    return wb


if __name__ == "__main__":
    current = read_current_stock()
    purchases, last_sale, dates, counts = read_vouchers()
    allocations, unmatched = allocate_fifo(current, purchases, last_sale)
    wb = build_workbook(current, allocations, unmatched, dates, counts, last_sale)
    wb.save(OUTPUT_FILE)
    print(
        json.dumps(
            {
                "output": str(OUTPUT_FILE),
                "stock_rows": len(current),
                "allocation_rows": len(allocations),
                "unmatched_rows": len(unmatched),
                "total_pcs": round(sum(r["qty_pcs"] for r in current), 3),
                "total_cases": round(sum(r["qty_cases"] for r in current), 4),
                "voucher_range": [min(dates).strftime("%Y-%m-%d") if dates else "", max(dates).strftime("%Y-%m-%d") if dates else ""],
            },
            indent=2,
        )
    )
